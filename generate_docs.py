import os
import math
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule

base_dir = r"C:\Users\erolt\Desktop\projekt_handlungsaufgabe\docs"
calc_dir = os.path.join(base_dir, "Berechnungen")
org_dir = os.path.join(base_dir, "01_Projektorganisation")

os.makedirs(calc_dir, exist_ok=True)
os.makedirs(org_dir, exist_ok=True)

# Define styles
bold_font = Font(bold=True)
header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
highlight_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")

thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

def apply_header_style(ws, row, cols):
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = bold_font
        cell.fill = header_fill
        cell.border = thin_border

# --- AUFGABE B: Kundenberechnungen ---
wb_calc = Workbook()
ws_speicher = wb_calc.active
ws_speicher.title = "Speicherplatzbedarf"

headers_speicher = ["Beschreibung", "Wert", "Einheit"]
ws_speicher.append(headers_speicher)
apply_header_style(ws_speicher, 1, 3)

data_speicher = [
    ["Anzahl Arbeitsplätze", 16, ""],
    ["Daten pro Tag/Arbeitsplatz", 350, "MiB"],
    ["Arbeitstage", 65, "Tage"],
    ["Gesamtdaten (MiB)", "=B2*B3*B4", "MiB"],
    ["Gesamtdaten (GiB)", "=B5/1024", "GiB"],
    ["Aufgerundet (GiB)", "=ROUNDUP(B6, 0)", "GiB"],
    ["Sicherheitszuschlag", "25%", ""],
    ["Gesamtbedarf inkl. Zuschlag (GiB)", "=ROUNDUP(B7*1.25, 0)", "GiB"]
]

for row in data_speicher:
    ws_speicher.append(row)

ws_speicher.cell(row=9, column=2).font = bold_font
ws_speicher.cell(row=9, column=2).fill = highlight_fill

ws_speicher.column_dimensions['A'].width = 35
ws_speicher.column_dimensions['B'].width = 15
ws_speicher.column_dimensions['C'].width = 10

ws_energie = wb_calc.create_sheet(title="Energiekosten")
headers_energie = ["Arbeitsplatztyp", "Anzahl", "Leistung/PC (W)", "Gesamtleistung (W)", "Arbeitstage", "Stunden/Tag", "Energie (kWh/Jahr)", "Strompreis (€/kWh)", "Kosten (€/Jahr)"]
ws_energie.append(headers_energie)
apply_header_style(ws_energie, 1, 9)

data_energie = [
    ["Standard-PCs", 12, 140, "=B2*C2", 220, 8.5, "=D2*E2*F2/1000", 0.32, "=G2*H2"],
    ["CAD-Workstations", 4, 320, "=B3*C3", 220, 8.5, "=D3*E3*F3/1000", 0.32, "=G3*H3"],
    ["Gesamt", "", "", "=SUM(D2:D3)", "", "", "=SUM(G2:G3)", "", "=SUM(I2:I3)"]
]

for row in data_energie:
    ws_energie.append(row)

ws_energie.cell(row=4, column=9).font = bold_font
ws_energie.cell(row=4, column=9).fill = highlight_fill
ws_energie.cell(row=4, column=1).font = bold_font

for col in ['A', 'G', 'I']:
    ws_energie.column_dimensions[col].width = 20

wb_calc.save(os.path.join(calc_dir, "Kundenberechnungen.xlsx"))


# --- AUFGABE 1a: Vorgangsliste ---
wb_vorgang = Workbook()
ws_vorgang = wb_vorgang.active
ws_vorgang.title = "Vorgangsliste"

headers_vorgang = ["Vorgangs-Nr", "Vorgangsbezeichnung", "Vorgänger", "Nachfolger", "Dauer (Tage)", "Verantwortlich"]
ws_vorgang.append(headers_vorgang)
apply_header_style(ws_vorgang, 1, 6)

data_vorgang = [
    [1, "Projektstart & Kick-off", "-", 2, 2, "Jan Kluth"],
    [2, "Anforderungsanalyse & Lastenheft", 1, 3, 5, "Mathias Vonau"],
    [3, "Marktrecherche & Produktauswahl", 2, 4, 5, "Marian Bolecke"],
    [4, "Nutzwertanalyse", 3, 5, 3, "Marian Bolecke"],
    [5, "Angebotserstellung & Kalkulation", 4, 6, 4, "Marian Bolecke"],
    [6, "Kundenfreigabe abwarten", 5, 7, 3, "Jan Kluth"],
    [7, "Bestellung & Beschaffung", 6, 8, 10, "Marian Bolecke"],
    [8, "Wareneingang & Prüfung", 7, 9, 2, "Marco Schmidt"],
    [9, "Vorinstallation & Konfiguration", 8, "10, 11", 5, "Mathias Vonau, Siyar"],
    [10, "Entsorgung Altdrucker", 9, 12, 2, "Marco Schmidt"],
    [11, "Lieferung & Montage vor Ort", 9, 12, 3, "Mathias Vonau, Siyar"],
    [12, "Funktionstest & Abnahme", "10, 11", 13, 2, "gesamtes Team"],
    [13, "Übergabe & Dokumentation", 12, 14, 1, "Jan Kluth"],
    [14, "Projektabschluss & Lessons Learned", 13, "-", 1, "gesamtes Team"]
]

for row in data_vorgang:
    ws_vorgang.append(row)

ws_vorgang.column_dimensions['B'].width = 40
ws_vorgang.column_dimensions['F'].width = 25
wb_vorgang.save(os.path.join(org_dir, "Vorgangsliste.xlsx"))

# --- AUFGABE 1b: Gantt-Diagramm ---
wb_gantt = Workbook()
ws_gantt = wb_gantt.active
ws_gantt.title = "Gantt-Diagramm"

ws_gantt.cell(row=1, column=1, value="Vorgang")
ws_gantt.column_dimensions['A'].width = 35

colors = {
    "Jan Kluth": "4F81BD",
    "Mathias Vonau": "C0504D",
    "Marian Bolecke": "9BBB59",
    "Marco Schmidt": "8064A2",
    "Mathias Vonau, Siyar": "F79646",
    "gesamtes Team": "4BACC6"
}

for day in range(1, 49):
    ws_gantt.cell(row=1, column=day+1, value=day)
    ws_gantt.column_dimensions[openpyxl.utils.get_column_letter(day+1)].width = 3

start_days = {}

for idx, row in enumerate(data_vorgang):
    v_id = row[0]
    v_name = row[1]
    v_pre = row[2]
    v_dur = row[4]
    v_resp = row[5]
    
    ws_gantt.cell(row=idx+2, column=1, value=f"{v_id}. {v_name}")
    
    if v_pre == "-":
        start_day = 1
    elif isinstance(v_pre, int):
        start_day = start_days[v_pre]
    else:
        pres = [int(p.strip()) for p in str(v_pre).split(",")]
        start_day = max(start_days[p] for p in pres)
        
    start_days[v_id] = start_day + v_dur
    
    fill = PatternFill(start_color=colors.get(v_resp, "000000"), end_color=colors.get(v_resp, "000000"), fill_type="solid")
    for d in range(start_day, start_day + v_dur):
        ws_gantt.cell(row=idx+2, column=d+1).fill = fill

# Legende
leg_row = len(data_vorgang) + 4
ws_gantt.cell(row=leg_row, column=1, value="Legende (Verantwortlichkeiten)").font = bold_font
for i, (resp, color) in enumerate(colors.items()):
    ws_gantt.cell(row=leg_row + i + 1, column=1, value=resp)
    ws_gantt.cell(row=leg_row + i + 1, column=2).fill = PatternFill(start_color=color, end_color=color, fill_type="solid")

wb_gantt.save(os.path.join(org_dir, "Gantt_Diagramm.xlsx"))

# --- AUFGABE 1c: Risikoanalyse ---
wb_risiko = Workbook()
ws_risiko = wb_risiko.active
ws_risiko.title = "Risikoanalyse"

headers_risiko = ["Risiko-Nr", "Risikobeschreibung", "Eintrittswahrscheinlichkeit (1-5)", "Schadensausmaß (1-5)", "Risikobewertung", "Gegenmaßnahmen"]
ws_risiko.append(headers_risiko)
apply_header_style(ws_risiko, 1, 6)

data_risiko = [
    [1, "Lieferverzögerung bei Hardware", 3, 4, "=C2*D2", "Ausweichlieferanten definieren, frühzeitige Bestellung"],
    [2, "Defekte Ware bei Lieferung", 2, 3, "=C3*D3", "Wareneingangsprüfung sofort durchführen, Puffer einplanen"],
    [3, "Inkompatibilität Software/Hardware", 2, 4, "=C4*D4", "Vorab-Kompatibilitätsprüfung, Testsystem aufbauen"],
    [4, "Budgetüberschreitung", 2, 5, "=C5*D5", "Regelmäßiges Kostencontrolling, Pufferbudget"],
    [5, "Personalausfall im Projektteam", 2, 3, "=C6*D6", "Wissensmanagement, Stellvertreterregelungen"],
    [6, "Netzwerkprobleme beim Rollout", 3, 3, "=C7*D7", "Vorab-Netzwerkanalyse, Backup-Lösungen"],
    [7, "Fehlende Kundenfreigabe verzögert Projekt", 3, 4, "=C8*D8", "Enger Kundenkontakt, klare Meilensteine kommunizieren"]
]

for row in data_risiko:
    ws_risiko.append(row)

ws_risiko.column_dimensions['B'].width = 40
ws_risiko.column_dimensions['C'].width = 30
ws_risiko.column_dimensions['D'].width = 25
ws_risiko.column_dimensions['E'].width = 15
ws_risiko.column_dimensions['F'].width = 50

red_fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
green_fill = PatternFill(start_color="00FF00", end_color="00FF00", fill_type="solid")

ws_risiko.conditional_formatting.add('E2:E8', CellIsRule(operator='greaterThanOrEqual', formula=['12'], stopIfTrue=True, fill=red_fill))
ws_risiko.conditional_formatting.add('E2:E8', CellIsRule(operator='between', formula=['8', '11'], stopIfTrue=True, fill=yellow_fill))
ws_risiko.conditional_formatting.add('E2:E8', CellIsRule(operator='lessThan', formula=['8'], stopIfTrue=True, fill=green_fill))

wb_risiko.save(os.path.join(org_dir, "Risikoanalyse.xlsx"))

# --- AUFGABE 1d: Projektstrukturplan.md ---
psp_content = """# Projektstrukturplan

```mermaid
graph TD
    A[Projekt VektorPlan IT-Ausstattung]
    
    A --> B[Hardware]
    A --> C[Software]
    A --> D[Netzwerk]
    A --> E[Services]
    A --> F[Projektmanagement]
    
    B --> B1[PCs]
    B --> B2[Workstations]
    B --> B3[Monitore]
    B --> B4[Peripherie]
    B --> B5[MFP]
    
    C --> C1[OS]
    C --> C2[Office]
    C --> C3[PDF]
    C --> C4[Security]
    
    D --> D1[Patchkabel]
    D --> D2[Konfiguration]
    
    E --> E1[Beratung]
    E --> E2[Installation]
    E --> E3[Rollout]
    E --> E4[Schulung]
    
    F --> F1[Planung]
    F --> F2[Controlling]
    F --> F3[Dokumentation]
    F --> F4[Abschluss]
```
"""

with open(os.path.join(org_dir, "Projektstrukturplan.md"), "w", encoding="utf-8") as f:
    f.write(psp_content)

print("All documents generated successfully.")
