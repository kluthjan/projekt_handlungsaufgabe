import os
import subprocess
import sys

def install_and_import():
    try:
        import docx
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
        import docx
install_and_import()

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# --- Paths ---
out_dir = r"C:\Users\erolt\Desktop\projekt_handlungsaufgabe\docs\07_Rollout"
os.makedirs(out_dir, exist_ok=True)
checklist_path = os.path.join(out_dir, "Checkliste_Inbetriebnahme.docx")
protocol_path = os.path.join(out_dir, "Abnahmeprotokoll.docx")

def create_checklist():
    doc = Document()
    
    # Title
    title = doc.add_heading('Checkliste Inbetriebnahme', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Info Text
    doc.add_paragraph('Projekt: 16 neue Arbeitsplätze (VektorPlan GmbH)')
    
    # Sections
    sections = {
        '1. Vorbereitung': [
            'Lieferung vollständig geprüft',
            'Arbeitsplatz vorbereitet und gereinigt',
            'Netzwerkdosen geprüft',
            'Werkzeug und Material bereitgestellt'
        ],
        '2. Hardware-Aufbau': [
            'PC/Workstation aufgestellt (ergonomische Position)',
            'Monitor 1 montiert (höhenverstellbar eingestellt)',
            'Monitor 2 montiert (höhenverstellbar eingestellt)',
            'Monitorhöhe: Oberkante auf Augenhöhe des Nutzers',
            'Tastatur und Maus platziert (Unterarme waagerecht)',
            'Kabelmanagement: Kabel gebündelt und beschriftet',
            'Patchkabel angeschlossen (Port dokumentiert)',
            'Externe SSD angeschlossen'
        ],
        '3. Ergonomie (nach ArbStättV & Bildschirmarbeitsverordnung)': [
            'Sehabstand zum Monitor: 50-70 cm',
            'Blickwinkel: leicht nach unten geneigt',
            'Keine Blendung/Reflexionen auf dem Bildschirm',
            'Ausreichende Beleuchtung (500 Lux)',
            'Stuhl ergonomisch eingestellt'
        ],
        '4. Software & Konfiguration': [
            'Windows 11 Pro aktiviert und lizenziert',
            'Windows Updates installiert',
            'Computername nach Schema vergeben (z.B. VP-PC-01)',
            'Domäne/Arbeitsgruppe konfiguriert',
            'Microsoft 365 installiert und lizenziert',
            'PDF-Software installiert',
            'Endpoint Protection installiert und aktuell',
            'Netzlaufwerke verbunden (Fileserver)',
            'Drucker eingerichtet (MFP)',
            'Energiespareinstellungen konfiguriert'
        ],
        '5. Funktionstest': [
            'Boot-Test erfolgreich',
            'Beide Monitore erkannt und korrekt konfiguriert',
            'Netzwerkverbindung aktiv (Ping Fileserver)',
            'Internet-Zugang funktioniert',
            'Druckertest-Seite erfolgreich',
            'USB-Geräte funktionieren',
            'Tastatur/Maus-Set verbunden und funktional',
            'Externe SSD erkannt und beschreibbar',
            '(CAD-Workstation) GPU-Treiber installiert und funktional',
            '(CAD-Workstation) Benchmark/Stresstest bestanden'
        ],
        '6. Abschluss': [
            'Arbeitsplatz-Dokumentation ausgefüllt',
            'Seriennummern erfasst',
            'Mitarbeiter eingewiesen'
        ]
    }
    
    for section, items in sections.items():
        doc.add_heading(section, level=1)
        for item in items:
            p = doc.add_paragraph()
            run = p.add_run('☐  ' + item)
            
    doc.add_paragraph("\n")
    
    # Signature table
    table = doc.add_table(rows=2, cols=4)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Arbeitsplatz-Nr.'
    hdr_cells[1].text = 'Datum'
    hdr_cells[2].text = 'Techniker'
    hdr_cells[3].text = 'Unterschrift'
    
    for i in range(4):
        table.rows[1].cells[i].text = ' '
        
    doc.save(checklist_path)


def create_protocol():
    doc = Document()
    
    # Title
    title = doc.add_heading('Abnahme- und Übergabeprotokoll', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Kopfbereich
    doc.add_heading('Projektdaten', level=1)
    
    table = doc.add_table(rows=5, cols=2)
    table.autofit = True
    
    data = [
        ('Projekt:', 'IT-Ausstattung Planungs- und Konstruktionsbüro'),
        ('Auftraggeber:', 'VektorPlan GmbH, [Adresse]'),
        ('Auftragnehmer:', 'NextLevel IT Solutions GmbH, [Adresse]'),
        ('Auftragsnummer:', 'NL-2026-VP-001'),
        ('Datum der Abnahme / Ort:', '____________________ / VektorPlan GmbH, Großraumbüro 1 & 2')
    ]
    
    for i, (k, v) in enumerate(data):
        row = table.rows[i].cells
        row[0].text = k
        row[1].text = v
        row[0].paragraphs[0].runs[0].bold = True
        
    # Leistungsübersicht
    doc.add_heading('Leistungsübersicht', level=1)
    
    items = [
        ('1', 'Standard-Büro-PC', '12', '12', '☐ OK  ☐ Mangel'),
        ('2', 'CAD-Workstation', '4', '4', '☐ OK  ☐ Mangel'),
        ('3', 'Monitor 27" WQHD', '32', '32', '☐ OK  ☐ Mangel'),
        ('4', 'Tastatur/Maus-Set', '16', '16', '☐ OK  ☐ Mangel'),
        ('5', 'Externe SSD (1TB)', '16', '16', '☐ OK  ☐ Mangel'),
    ]
    
    table_items = doc.add_table(rows=1, cols=5)
    table_items.style = 'Table Grid'
    hdr = table_items.rows[0].cells
    for i, h in enumerate(['Pos.', 'Bezeichnung', 'Menge Soll', 'Menge Ist', 'Status']):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True
        
    for item in items:
        row = table_items.add_row().cells
        for i, val in enumerate(item):
            row[i].text = val

    # Funktionstest-Ergebnis
    doc.add_heading('Funktionstest-Ergebnis', level=1)
    tests = [
        'Alle Arbeitsplätze getestet',
        'Alle Softwarelizenzen aktiviert',
        'Netzwerkanbindung funktional',
        'Drucker funktional'
    ]
    for t in tests:
        doc.add_paragraph(f'{t}: ☐ Ja  ☐ Nein')

    # Mängelliste
    doc.add_heading('Mängelliste', level=1)
    table_m = doc.add_table(rows=2, cols=4)
    table_m.style = 'Table Grid'
    hdr = table_m.rows[0].cells
    headers = ['Nr', 'Beschreibung', 'Kategorie (wesentlich/unwesentlich)', 'Behebungsfrist']
    for i, h in enumerate(headers):
        hdr[i].text = h
        hdr[i].paragraphs[0].runs[0].bold = True
        
    # Abnahmeerklärung
    doc.add_heading('Abnahmeerklärung', level=1)
    doc.add_paragraph('Die oben aufgeführten Leistungen wurden geprüft und werden hiermit [vorbehaltlos / unter Vorbehalt der Beseitigung der aufgeführten Mängel] abgenommen.')

    # Unterschriften
    doc.add_paragraph('\n')
    table_sig = doc.add_table(rows=2, cols=2)
    sig_data = [
        ('Auftraggeber (Name, Datum, Unterschrift)', 'Auftragnehmer (Name, Datum, Unterschrift)'),
        ('\n\n_________________________________', '\n\n_________________________________')
    ]
    for i, row_data in enumerate(sig_data):
        row = table_sig.rows[i].cells
        row[0].text = row_data[0]
        row[1].text = row_data[1]

    # Anlagen
    doc.add_heading('Anlagen', level=1)
    doc.add_paragraph('☐ Seriennummern-Liste\n☐ Lizenznachweise\n☐ Garantieunterlagen')

    doc.save(protocol_path)


if __name__ == '__main__':
    create_checklist()
    create_protocol()
    print("Documents created successfully.")
